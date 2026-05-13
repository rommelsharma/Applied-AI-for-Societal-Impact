#!/usr/bin/env python3
"""Regenerate ``docs/Solution Design Document.docx``.

The body is ordered to match ``docs/rag_solution_design_best_practices.docx``
(playbook sections 1–15), then delivery sequence, pitfalls, and appendices.
Requires: ``python-docx`` (install via ``pip install python-docx``).
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "docs" / "Solution Design Document.docx"


def _bullets(doc, items: list[str]) -> None:
    for line in items:
        doc.add_paragraph(line, style="List Bullet")


def main() -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:  # pragma: no cover
        print("Install python-docx: pip install python-docx", file=sys.stderr)
        raise SystemExit(1) from exc

    doc = Document()

    title = doc.add_heading("Solution Design Document", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.add_run("Bias-Aware Decision Making RAG Copilot\n").bold = True
    st.add_run("Decision Intelligence | Applied AI for Societal Impact\n\n")
    st.add_run("Rommel Sharma  |  Version 1.1  |  Issue date: 13 May 2026\n")
    st.add_run("Classification: Portfolio / Public\n")

    doc.add_paragraph()
    doc.add_heading("Document control", level=1)
    doc.add_paragraph(
        "This revision restructures the solution narrative to follow the same lifecycle sequence as "
        "``docs/rag_solution_design_best_practices.docx`` (synthesised guidance from Google Cloud, "
        "Microsoft Azure Architecture Center, Redwerk, and ChatRAG). Technical content from prior "
        "versions is preserved but reorganised so reviewers can map each engineering decision to an "
        "industry-recognised practice."
    )
    _bullets(
        doc,
        [
            "v0.1 — Initial scaffolding.",
            "v1.0 — 08 May 2026 — Baseline portfolio design.",
            "v1.1 — 13 May 2026 — Playbook-aligned sequence; adds eval spine (metrics, connectivity, "
            "baseline scenario capture, run cards) and explicit gap/roadmap call-outs.",
        ],
    )

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "The Bias-Aware Decision Making RAG Copilot is decision-support software (not legal, medical, "
        "HR-policy, or financial advice) that analyses free-text scenarios for cognitive and systemic "
        "biases. It combines (1) a curated decision-science corpus, (2) Amazon Bedrock (Titan Text "
        "Embeddings v2 + Claude Sonnet 4.5) with FAISS vector retrieval, optional MMR diversification, "
        "optional Claude-Haiku reranking, and concept/decision-domain filters, and (3) a closed "
        "bias taxonomy embedded in a strict JSON output contract. Every scenario is executed twice—"
        "baseline (no retrieval) and RAG-enhanced—so lift from retrieval is measurable, not assumed."
    )
    doc.add_paragraph(
        "Recent engineering adds a reproducible evaluation spine: pre-flight Bedrock + vector-store "
        "connectivity checks, per-response schema/taxonomy metrics, aggregate rollups, and run cards "
        "(git commit, prompt/scenario hashes, model and RAG settings). Timestamped captures under "
        "``response/`` (via ``scripts/record_response_run.py``) anchor before/after comparisons for "
        "the same three frozen scenarios in ``evaluation/baseline_scenarios.json``."
    )

    doc.add_heading("System architecture snapshot (cross-reference)", level=1)
    doc.add_paragraph(
        "Offline chain (corpus refresh): raw PDFs → ``pdf_parser.py`` (parsed text + ``.meta.json``) → "
        "``chunker.py`` (``chunks.json``) → ``concept_extractor.py`` (``chunks_with_concepts.json``) → "
        "``enrich_chunks.py`` (``knowledge_base.json``) → ``build_vector_index.py`` "
        "(``embeddings.npy``, ``knowledge.index``, aligned ``index_metadata.json``, ``manifest.json``)."
    )
    doc.add_paragraph(
        "Runtime chain: user scenario → ``extract_concepts`` (taxonomy) → ``KnowledgeRetriever.search`` "
        "(Titan query embedding, FAISS candidates, optional filters + optional Haiku rerank + MMR) → "
        "``format_retrieved_context`` → two ``BedrockProvider.converse`` calls (baseline without context "
        "blocks, RAG with context + same taxonomy appendix) → JSON payloads plus retrieval trace. "
        "For file-level detail see ``docs/code-flow.md``."
    )

    # --- Playbook sections 1–15 ---
    doc.add_heading(
        "1. Define the use case and success criteria (playbook §1)",
        level=1,
    )
    doc.add_paragraph(
        "Use case: structured second-opinion analysis when stakes are high—hiring and promotion "
        "committees, performance calibration, AI governance reviews, compliance pattern review, "
        "and leadership trade-offs. Personas: managers, HR/People Ops, risk and compliance, legal "
        "reviewers, and portfolio assessors."
    )
    doc.add_paragraph(
        "Measurable pilot signals: JSON schema validity; bias names constrained to ``data/metadata/"
        "bias-taxonomy.json``; retrieval trace present; baseline vs RAG output pairs persisted; "
        "``evaluation/metrics.py`` flags taxonomy violations; ``scripts/record_response_run.py`` "
        "produces comparable artefacts over time. Non-functional: auditability (chunk → source), "
        "low-temperature reproducibility (0.1), AWS-native posture (Bedrock, IAM, future CloudWatch)."
    )
    doc.add_paragraph(
        "Gap / next: formal latency and cost SLO dashboards; richer pilot KPIs (groundedness judges, "
        "human rubrics) as in playbook §14."
    )

    doc.add_heading("2. Prepare the knowledge foundation (playbook §2)", level=1)
    doc.add_paragraph(
        "Corpus under ``data/corpora/<corpus>/`` with public synthesised dossiers (tracked PDFs) and "
        "optional private full-book material (gitignored). Ingestion is trustworthy-by-construction "
        "for the public path: project-authored dossiers only. ``data_pipeline/pdf_parser.py`` extracts "
        "text and sidecar metadata; ``chunker.py`` removes front/back matter noise."
    )
    doc.add_paragraph(
        "Gap / next: explicit source-system inventory for enterprise expansion; automated deduplication, "
        "staleness, and supersession workflow; event-driven re-embed/reindex when raw files change."
    )

    doc.add_heading("3. Build a representative test set early (playbook §3)", level=1)
    doc.add_paragraph(
        "``evaluation/test_scenarios.json`` holds a broad scenario suite. ``evaluation/baseline_scenarios.json`` "
        "freezes three diverse scenarios (hiring rubric gap, AI governance / proxy risk, strategy sunk-cost) "
        "for longitudinal comparison. ``evaluation/run_evaluation.py`` runs the full suite to "
        "``evaluation/latest_results.json``."
    )
    doc.add_paragraph(
        "``scripts/record_response_run.py`` runs connectivity (credentials → embedding probe → vector load → "
        "retrieval smoke), executes each baseline scenario through ``detect_bias_comparison``, attaches "
        "``evaluation/metrics.py`` scores, embeds a ``evaluation/run_card.py`` snapshot, and writes "
        "``response/<timestamp>_<label>_rag_eval.json``."
    )
    doc.add_paragraph(
        "Gap / next: golden reference answers; explicit out-of-corpus decline scenarios; typo and "
        "multi-part stress variants; automated retrieval-precision metrics against labelled chunk ids."
    )

    doc.add_heading("4. Analyse documents before chunking (playbook §4)", level=1)
    doc.add_paragraph(
        "``pdf_parser.py`` drives structure awareness before split: regex heading cues for chapter "
        "starts, ``parser_notes`` when metadata is inferred. ``chunker.py`` applies line-anchored "
        "patterns to strip prefaces, indexes, bibliographies, and similar boilerplate so chunks embed "
        "substantive argument rather than navigation text."
    )
    doc.add_paragraph(
        "Gap / next: optional markdown-normalisation path for complex PDF layouts when the corpus "
        "expands beyond clean dossiers."
    )

    doc.add_heading("5. Structure documents for retrieval quality (playbook §5)", level=1)
    doc.add_paragraph(
        "Dossiers are authored with descriptive titles and tight topical focus. ``enrich_chunks.py`` "
        "adds summaries, keywords, ``decision_domains``, and passage classification so each row in "
        "``knowledge_base.json`` is more self-contained when retrieved out of order."
    )

    doc.add_heading("6. Choose a chunking strategy deliberately (playbook §6)", level=1)
    doc.add_paragraph(
        "``data_pipeline/chunker.py`` uses profile-driven overlapping windows: ``dossier`` (350 words / "
        "60 overlap), ``book`` (500 / 80), ``auto`` selecting ``book`` when the parsed work exceeds "
        "~30k words. Chapter titles propagate to chunks. Strategy is documented in-module; corpus "
        "operators choose ``CHUNK_PROFILE`` via environment."
    )
    doc.add_paragraph(
        "Gap / next: automated experiment grid (profiles × frozen eval) to optimise recall vs noise, "
        "plus optional LLM-assisted chunk anchors for very long sections."
    )

    doc.add_heading("7. Clean and enrich chunks with metadata (playbook §7)", level=1)
    doc.add_paragraph(
        "Whitespace normalisation, cleaning notes, stable ``id``, ``source``, ``author``, ``title``, "
        "``publication_year``, ``chapter_title``, ``chunking_profile``, ``concepts``, ``concept_confidence``, "
        "``importance``, ``decision_domains``, ``keywords``, ``passage_type``, ``decision_phase``, "
        "``summary``, and ``text``—all carried into ``index_metadata.json`` in strict row alignment with "
        "``embeddings.npy`` / FAISS."
    )
    doc.add_paragraph(
        "Gap / next: enterprise metadata (source system, effective date, geography, product) and "
        "permission / ACL tags for filtered retrieval."
    )

    doc.add_heading("8. Select and evaluate the embedding model (playbook §8)", level=1)
    doc.add_paragraph(
        "Amazon Titan Text Embeddings v2 at 1024 dimensions with ``normalize=true``. The same model "
        "embeds corpus chunks at index time and user scenarios at query time via ``BedrockProvider."
        "embed_text()``—no asymmetric dual-encoder unless deliberately introduced later."
    )
    doc.add_paragraph(
        "Rationale: managed Bedrock surface, inner-product on unit vectors equals cosine similarity "
        "matching ``IndexFlatIP``. Gap / next: benchmark an alternate embedding on held-out queries "
        "with identical eval harness before switching ``BEDROCK_EMBEDDING_MODEL_ID``."
    )

    doc.add_heading("9. Design the index and retrieval strategy (playbook §9)", level=1)
    doc.add_paragraph(
        "Artefacts: ``embeddings.npy``, ``knowledge.index`` (FAISS ``IndexFlatIP``), ``index_metadata."
        "json``, ``manifest.json``. ``rag/retriever.py`` embeds the query, pulls an expanded candidate pool, "
        "applies concept and ``decision_domains`` filters with an unfiltered fallback, optionally reranks "
        "with Claude Haiku (``RAG_RERANKER``), then applies MMR diversification so one chapter cannot "
        "dominate the context window."
    )
    doc.add_paragraph(
        "Gap / next: hybrid vector + lexical (BM25) fusion; query decomposition for multi-part "
        "scenarios; metadata security filters once ACL fields exist."
    )

    doc.add_heading("10. Re-rank and assemble context carefully (playbook §10)", level=1)
    doc.add_paragraph(
        "Optional LLM reranker re-orders the candidate list; MMR produces the final cut. "
        "``format_retrieved_context`` in ``bias_detector.py`` emits numbered blocks with source, author, "
        "chunk id, passage type, decision phase, concepts, summary, and a bounded text excerpt for the "
        "generator."
    )
    doc.add_paragraph(
        "Gap / next: token-budget-aware assembly modes (excerpt vs full chunk vs expanded section) "
        "selected by experiment, not fixed heuristics only."
    )

    doc.add_heading("11. Design prompts and generation for grounded answers (playbook §11)", level=1)
    doc.add_paragraph(
        "``prompts/bias_detection_system_prompt.txt`` locks role, disclaimer, taxonomy-only bias names, "
        "uncertainty when context is weak, and JSON fields: ``situation_summary``, ``biases_identified``, "
        "``noise_considerations``, ``overall_risk_assessment``, ``recommended_actions``, "
        "``alternative_perspectives``. Chat calls use temperature 0.1 for repeatability."
    )
    doc.add_paragraph(
        "Gap / next: versioned prompt files (v2, v3) with automated A/B harness tied to the same "
        "``baseline_scenarios.json`` payloads."
    )

    doc.add_heading(
        "12. Evaluate retrieval and generation separately, then end-to-end (playbook §12)",
        level=1,
    )
    doc.add_paragraph(
        "Connectivity isolates Bedrock and vector-store health. ``evaluation/metrics.py`` scores each "
        "model path for required JSON shape and taxonomy membership before comparing baseline vs RAG "
        "bias counts and retrieval score summaries. End-to-end payloads remain the source of truth for "
        "qualitative review."
    )
    doc.add_paragraph(
        "Gap / next: labelled gold chunks per scenario for retrieval hit-rate; LLM-as-judge and "
        "reference-based scores (where references exist); CI regression against a frozen "
        "``baseline_metrics.json``."
    )

    doc.add_heading("13. Controlled experimentation and root-cause analysis (playbook §13)", level=1)
    doc.add_paragraph(
        "``evaluation/run_card.py`` records UTC and local timestamps, git commit and dirty flag, "
        "SHA-256 of the scenario file and system prompt, Bedrock model ids, embedding dimensions, "
        "corpus name, chunk profile, ``RAG_TOP_K``, MMR lambda, reranker mode and candidate depth. "
        "Paired with label-based filenames under ``response/``, teams can change one knob at a time "
        "and keep an audit trail."
    )

    doc.add_heading("14. Human evaluation before production rollout (playbook §14)", level=1)
    doc.add_paragraph(
        "Automated checks are necessary but not sufficient. ``docs/sample_results_comparison.md`` "
        "supports qualitative review. Planned extension: structured reviewer rubrics (clarity, "
        "trustworthiness, citation usefulness) feeding back into corpus and prompt workstreams."
    )

    doc.add_heading("15. Operations, governance, and continuous improvement (playbook §15)", level=1)
    doc.add_paragraph(
        "``docs/ml-ops.plan.md`` remains the operational north star: versioned corpora, manifests, "
        "CI smoke paths, drift and cost thinking. Runtime logging of token usage and per-step latency "
        "is the next hook once a deployed API exists. Re-evaluation triggers when taxonomy, corpus, "
        "embedding model, or retrieval logic changes materially."
    )

    doc.add_heading("Recommended delivery sequence (playbook synthesis)", level=1)
    _bullets(
        doc,
        [
            "Define use case, personas, risks, and pilot metrics.",
            "Curate and govern source content (public dossiers today; private books under strict controls).",
            "Build and freeze evaluation scenarios early; add gold labels incrementally.",
            "Analyse document structures; extend parsers as new formats appear.",
            "Tune chunking with measured retrieval and answer quality—not intuition alone.",
            "Enrich metadata for filters, citations, and future ACLs.",
            "Benchmark embeddings when corpus or question style shifts.",
            "Design retrieval logic (vector today; hybrid and decomposition on roadmap).",
            "Add reranking and context-assembly experiments.",
            "Iterate prompts with version control and A/B metrics.",
            "Automate retrieval/generation/E2E evaluation with regression gates.",
            "Run human evaluation; deploy in phases; monitor continuously.",
        ],
    )

    doc.add_heading("Common pitfalls and mitigations", level=1)
    _bullets(
        doc,
        [
            "Low-quality sources — mitigated by curated dossiers and parser/chunk hygiene; risk grows "
            "if unvetted PDFs enter without review.",
            "Naive chunking — mitigated by overlap, chapter awareness, profiles, and planned experiment grid.",
            "Single retrieval mode — acknowledged gap; hybrid lexical + vector is planned.",
            "Prompt overload with weak chunks — mitigated by MMR, filters, reranker option; assembly modes "
            "still tunable.",
            "Optimising only end-to-end — mitigated by connectivity + schema/taxonomy metrics today; "
            "retrieval-precision metrics next.",
        ],
    )

    doc.add_heading("Appendix A — Runtime and offline module map", level=1)
    doc.add_paragraph(
        "Configuration: ``shared_components/settings.py`` (``BEDROCK_SETTINGS``, ``RAG_SETTINGS``). "
        "Paths: ``shared_components/utilities/path_utils.py``. Taxonomy helpers: "
        "``shared_components/utilities/taxonomy_utils.py``. Ingestion: ``data_pipeline/{pdf_parser,"
        "chunker,concept_extractor,enrich_chunks,build_knowledge_base,build_vector_index}.py``. "
        "Retrieval: ``rag/retriever.py``. Bedrock access: ``app/services/bedrock_provider.py``. "
        "Orchestration: ``app/services/bias_detector.py``. Evaluation: ``evaluation/*``; capture: "
        "``scripts/record_response_run.py``."
    )

    doc.add_heading("Appendix B — Reference hosting (summary)", level=1)
    doc.add_paragraph(
        "Target AWS layout unchanged in intent: CloudFront → API Gateway or ALB → Fargate/Lambda "
        "service → Bedrock (Claude + Titan); vector artefacts on S3; Secrets Manager for credentials; "
        "CloudWatch, X-Ray, CloudTrail for observability and audit. VPC endpoints for private Bedrock "
        "egress where required."
    )

    doc.add_heading("Appendix C — Roadmap alignment (engineering phases)", level=1)
    _bullets(
        doc,
        [
            "Phase 1 — Eval spine: ``evaluation/metrics.py`` (schema + taxonomy checks), extended scenarios, "
            "retrieval/generation metrics to grow over time.",
            "Phase 4 — Run cards: ``evaluation/run_card.py`` + timestamped ``response/*_rag_eval.json`` from "
            "``scripts/record_response_run.py`` for reproducible before/after comparisons.",
            "Phase 5 — Retrieval science: hybrid lexical + vector retrieval, query decomposition for "
            "multi-part scenarios, configurable context assembly modes.",
            "Phase 3 — ACL and audience-aware metadata on chunks (when multi-tenant or multi-audience "
            "deployments require entitlement filtering at retrieval time).",
            "Phase 2 — Content operations: source manifest, content hashes, incremental re-embed/reindex, "
            "deduplication and staleness handling.",
            "Phase 6 — Chunking grid automation plus expanded human evaluation and production telemetry "
            "(latency, tokens, failure rates).",
            "Phase 7 — Prompt versioning, optional explicit refusal / insufficient-evidence fields in JSON "
            "when retrieval is weak; prompt A/B harness tied to frozen scenarios.",
        ],
    )

    doc.add_paragraph()
    p = doc.add_paragraph(
        "Disclaimer: This system provides research-informed decision support only. It does not replace "
        "professional judgment or organisational policy."
    )
    for run in p.runs:
        run.italic = True

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
